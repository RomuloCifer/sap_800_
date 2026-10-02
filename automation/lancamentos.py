"""
Leitura da planilha de lançamentos (entrada/lancamentos.xlsx).

Usa:
  - Batch        → inicia a Parte 1
  - Total Value  → compara com os prints A e B (OCR)
  - Issue Date   → Parte 4: impostos extras CBS1/IB2S se ano >= 2026
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from automation.utils import clean_value
from automation.value_ocr import normalize_money

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_PLANILHA = ROOT / "entrada" / "lancamentos.xlsx"

# Nomes aceitos (case-insensitive, espaços normalizados)
_BATCH_ALIASES = {"batch", "batch number", "batch nr", "batch no"}
_TOTAL_ALIASES = {
    "total value",
    "totalvalue",
    "valor total",
    "total",
}
_ISSUE_DATE_ALIASES = {
    "issue date",
    "issuedate",
    "data emissao",
    "data emissão",
    "dt emissao",
    "dt emissão",
}


def _norm_header(value):
    # type: (object) -> str
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = " ".join(text.split())
    return text


def _cell_to_str(value):
    # type: (object) -> str
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return clean_value(str(value))


def parse_issue_year(value):
    # type: (object) -> Optional[int]
    """
    Extrai o ano de Issue Date (datetime Excel, str US/BR, etc.).
    Retorna None se vazio ou ilegível.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return int(value.year)
    if isinstance(value, date):
        return int(value.year)

    text = _cell_to_str(value)
    if not text:
        return None

    # ISO / Excel serial-ish: 2026-06-23
    m = re.match(r"^(\d{4})[-/.]", text)
    if m:
        return int(m.group(1))

    # Final com ano 4 dígitos: 6/23/2026, 23.06.2026, 23/06/2026
    m = re.search(r"(\d{4})\s*$", text)
    if m:
        return int(m.group(1))

    # Ano 2 dígitos no fim: 6/23/26 → 2026 (assume 2000+)
    m = re.search(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2})\s*$", text)
    if m:
        return 2000 + int(m.group(3))

    return None


def issue_date_is_2026_plus(value):
    # type: (object) -> bool
    year = parse_issue_year(value)
    if year is None:
        raise ValueError("Issue Date ilegível: {!r}".format(value))
    return year >= 2026


def _find_columns(headers):
    # type: (List[object]) -> Dict[str, int]
    """Retorna índices 0-based: batch, total_value, issue_date."""
    normalized = [_norm_header(h) for h in headers]
    batch_idx = None
    total_idx = None
    issue_idx = None
    for i, name in enumerate(normalized):
        if not name:
            continue
        if batch_idx is None and (name in _BATCH_ALIASES or name.startswith("batch")):
            batch_idx = i
        if total_idx is None and (
            name in _TOTAL_ALIASES or name.replace(" ", "") == "totalvalue"
        ):
            total_idx = i
        if issue_idx is None and (
            name in _ISSUE_DATE_ALIASES or name.replace(" ", "") == "issuedate"
        ):
            issue_idx = i
    missing = []
    if batch_idx is None:
        missing.append("Batch")
    if total_idx is None:
        missing.append("Total Value")
    if issue_idx is None:
        missing.append("Issue Date")
    if missing:
        raise ValueError(
            "Planilha sem coluna(s) obrigatória(s): {}. "
            "Cabeçalhos encontrados: {}".format(
                ", ".join(missing),
                [h for h in headers if h is not None],
            )
        )
    return {
        "batch": batch_idx,
        "total_value": total_idx,
        "issue_date": issue_idx,
    }


def load_lancamentos(path=None):
    # type: (Optional[object]) -> List[Dict[str, Any]]
    """
    Lê a planilha e devolve lista de dicts:
      {row, batch, total_value, total_value_raw, issue_date_raw,
       issue_year, taxes_2026_plus}

    Ignora linhas sem Batch.
    """
    path = Path(path) if path is not None else DEFAULT_PLANILHA
    if not path.is_file():
        raise FileNotFoundError(
            "Planilha não encontrada: {}. "
            "Coloque o arquivo em entrada/lancamentos.xlsx.".format(path)
        )

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise ValueError("Planilha vazia: {}".format(path))

        cols = _find_columns(list(header))
        out = []  # type: List[Dict[str, Any]]
        for excel_row, row in enumerate(rows, start=2):
            if row is None:
                continue
            cells = list(row)
            batch_raw = cells[cols["batch"]] if cols["batch"] < len(cells) else None
            total_raw = (
                cells[cols["total_value"]] if cols["total_value"] < len(cells) else None
            )
            issue_raw = (
                cells[cols["issue_date"]] if cols["issue_date"] < len(cells) else None
            )
            batch = _cell_to_str(batch_raw)
            if not batch:
                continue
            total_str = _cell_to_str(total_raw)
            total_dec = normalize_money(total_str) if total_str else None
            if total_dec is None and total_raw is not None:
                # Número Excel puro (float)
                try:
                    total_dec = normalize_money(str(total_raw))
                except Exception:
                    total_dec = None
            if total_dec is None:
                raise ValueError(
                    "Linha {}: Batch {!r} sem Total Value legível ({!r}).".format(
                        excel_row, batch, total_raw
                    )
                )
            issue_year = parse_issue_year(issue_raw)
            if issue_year is None:
                raise ValueError(
                    "Linha {}: Batch {!r} sem Issue Date legível ({!r}).".format(
                        excel_row, batch, issue_raw
                    )
                )
            out.append(
                {
                    "row": excel_row,
                    "batch": batch,
                    "total_value": total_dec,
                    "total_value_raw": total_str or str(total_raw),
                    "issue_date_raw": _cell_to_str(issue_raw) or str(issue_raw),
                    "issue_year": issue_year,
                    "taxes_2026_plus": issue_year >= 2026,
                }
            )
    finally:
        wb.close()

    if not out:
        raise ValueError("Nenhuma linha com Batch em {}.".format(path))

    print(
        "Planilha {}: {} lançamento(s) (Batch + Total Value + Issue Date).".format(
            path.name, len(out)
        )
    )
    for item in out:
        print(
            "  linha {}  batch={!r}  total_value={}  issue_date={!r}  "
            "ano={}  impostos_cbs_ibs={}".format(
                item["row"],
                item["batch"],
                item["total_value"],
                item["issue_date_raw"],
                item["issue_year"],
                "sim" if item["taxes_2026_plus"] else "nao",
            )
        )
    return out


def resolve_planilha_from_argv(argv=None):
    # type: (Optional[list]) -> Path
    """Aceita --planilha CAMINHO; senão usa entrada/lancamentos.xlsx."""
    args = list(argv if argv is not None else [])
    path = DEFAULT_PLANILHA
    if "--planilha" in args:
        i = args.index("--planilha")
        if i + 1 >= len(args):
            raise ValueError("Use: --planilha caminho\\para\\lancamentos.xlsx")
        path = Path(args[i + 1])
    return path
